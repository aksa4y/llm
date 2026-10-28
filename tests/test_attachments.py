import base64
import io
import shutil
import unittest
from unittest.mock import Mock
from attachments import extract, unpack_files, attach_context


class AttachmentTests(unittest.TestCase):
    def test_text_encodings_and_validation(self):
        for encoding in ('utf-8', 'utf-16', 'cp1251'):
            self.assertEqual(extract('текст.txt', 'Привет'.encode(encoding))['text'], 'Привет')
        with self.assertRaises(ValueError):
            extract('run.exe', b'anything')
        with self.assertRaises(ValueError):
            unpack_files([{'name':'file.txt','data':'not base64!'}])
        result = unpack_files([{'name':'../../file.txt','data':base64.b64encode(b'hello').decode()}])
        self.assertEqual(result[0]['name'], 'file.txt')

    def test_docx_and_xlsx(self):
        from docx import Document
        from openpyxl import Workbook
        document = Document(); document.add_paragraph('Русский документ')
        buffer = io.BytesIO(); document.save(buffer)
        self.assertIn('Русский документ', extract('test.docx', buffer.getvalue())['text'])
        book = Workbook(); book.active.append(['Товар', 'Цена']); book.active.append(['Чай', 100])
        buffer = io.BytesIO(); book.save(buffer)
        self.assertIn('Чай | 100', extract('test.xlsx', buffer.getvalue())['text'])

    def test_pdf_text_and_scan_detection(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter(); page = writer.add_blank_page(300, 300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 200 Td (HELLO PDF 42) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        buffer = io.BytesIO(); writer.write(buffer)
        self.assertIn('HELLO PDF 42', extract('test.pdf', buffer.getvalue())['text'])
        writer = PdfWriter(); writer.add_blank_page(300, 300)
        buffer = io.BytesIO(); writer.write(buffer)
        with self.assertRaisesRegex(ValueError, 'сканы'):
            extract('scan.pdf', buffer.getvalue())

    def test_attachment_context_and_visible_truncation(self):
        engine = Mock(); engine.count_tokens.side_effect = lambda messages: len(messages[-1]['content'])
        context, trimmed = attach_context('Объясни', [{'name':'note.txt','text':'Важная заметка'}], engine, 1536)
        self.assertIn('Важная заметка', context)
        self.assertFalse(trimmed)
        context, trimmed = attach_context('Объясни', [{'name':'big.txt','text':'А' * 30000}], engine, 3072)
        self.assertTrue(trimmed)
        self.assertLess(len(context) + 3072, 8192)

    @unittest.skipUnless(shutil.which('tesseract'), 'Tesseract is not installed')
    def test_real_image_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (900, 160), 'white')
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 50)
        ImageDraw.Draw(image).text((25, 35), 'HELLO DOCUMENT 42', font=font, fill='black')
        buffer = io.BytesIO(); image.save(buffer, format='PNG')
        self.assertIn('42', extract('test.png', buffer.getvalue())['text'])
